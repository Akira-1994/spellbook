"""Convert the reviewed taxonomy workbook into spellbook/taxonomy.json.

data/taxonomy/taxonomy-review.xlsx is the human-edited source of truth for
class names, level corrections and school overrides. Run this after editing it:

    .venv/Scripts/python.exe -m scripts.import_taxonomy
"""
from __future__ import annotations

import json
import re
import sqlite3
from contextlib import closing
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "data" / "taxonomy" / "taxonomy-review.xlsx"
SEED = ROOT / "data" / "spellbook.sqlite"
OUTPUT = ROOT / "spellbook" / "taxonomy.json"

KINDS = {"職業": "class", "領域": "domain", "其他": "other"}
SCHOOLS = [
    # key, name, pattern matched against the free-text school field. The PDF
    # translation renders Evocation as 招魂, Abjuration as 放棄 and
    # Enchantment as 魅力, so those mistranslations map to their real schools.
    ("abjuration", "防護", "防護|放棄"),
    ("conjuration", "咒法", "咒法|召喚"),
    ("divination", "預言", "預言|占卜"),
    ("enchantment", "惑控", "惑控|附魔|魅控|魅力"),
    ("evocation", "塑能", "塑能|招魂"),
    ("illusion", "幻術", "幻術|幻覺|錯覺"),
    ("necromancy", "死靈", "死靈"),
    ("transmutation", "變化", "變化|嬗變|變形|變身"),
    ("universal", "通用", "共通|通用"),
]


def split_names(value) -> list[str]:
    return [part.strip() for part in re.split("[；;]", str(value or "")) if part.strip()]


def rows(sheet):
    headers = [cell.value for cell in sheet[1]]
    for values in sheet.iter_rows(min_row=2, values_only=True):
        if any(value is not None for value in values):
            yield dict(zip(headers, values))


def main() -> None:
    book = load_workbook(WORKBOOK, read_only=True)
    school_names = {name: key for key, name, _ in SCHOOLS}

    classes = []
    for row in rows(book["標準職業表"]):
        kind = KINDS.get(row["類別"])
        if not kind:
            raise SystemExit(f"標準職業表「{row['標準名稱']}」的類別無效：{row['類別']}")
        classes.append({"name": row["標準名稱"].strip(), "kind": kind, "name_en": (row["英文名稱"] or "").strip()})
    known = {item["name"] for item in classes}

    aliases, fix_aliases, dropped = {}, [], []
    for row in rows(book["職業對照"]):
        raw, kind = row["原始寫法"], row["類別"]
        if kind == "等級修正":
            fix_aliases.append(raw)
        elif kind == "刪除":
            dropped.append(raw)
        elif kind in KINDS:
            names = split_names(row["標準名稱"])
            missing = [name for name in names if name not in known]
            if not names or missing:
                raise SystemExit(f"職業對照「{raw}」的標準名稱不在標準職業表：{missing or '（空白）'}")
            aliases[raw] = names
        else:
            raise SystemExit(f"職業對照「{raw}」的類別無效：{kind}")

    with closing(sqlite3.connect(f"file:{SEED.as_posix()}?mode=ro", uri=True)) as db:
        raw_names = {row[0] for row in db.execute("SELECT name FROM classes")}
        unmapped = raw_names - set(aliases) - set(fix_aliases) - set(dropped)
        if unmapped:
            raise SystemExit(f"職業對照缺少這些寫法：{sorted(unmapped)}")

        level_fixes = {}
        for row in rows(book["等級修正"]):
            entries = db.execute(
                "SELECT e.id FROM spells s JOIN spell_entries e ON e.spell_id=s.id WHERE s.name_zh=? AND s.name_en=?",
                (row["法術"], row["英文名稱"]),
            ).fetchall()
            if len(entries) != 1:
                raise SystemExit(f"等級修正「{row['法術']}」對應到 {len(entries)} 個條目")
            levels = []
            for part in split_names(row["建議正確等級"]):
                match = re.fullmatch(r"(.+?)\s*(\d)", part)
                if not match or match.group(1) not in known:
                    raise SystemExit(f"等級修正「{row['法術']}」無法解析：{part}")
                levels.append([match.group(1), int(match.group(2))])
            level_fixes[entries[0][0]] = levels
        uncovered = db.execute(
            f"""SELECT DISTINCT s.name_zh FROM spell_levels l JOIN classes c ON c.id=l.class_id
                JOIN spell_entries e ON e.id=l.spell_entry_id JOIN spells s ON s.id=e.spell_id
                WHERE c.name IN ({','.join('?' * len(fix_aliases))})""",
            fix_aliases,
        ).fetchall()
        not_fixed = [name for (name,) in uncovered if not any(
            db.execute("SELECT 1 FROM spell_entries e JOIN spells s ON s.id=e.spell_id WHERE e.id=? AND s.name_zh=?", (entry, name)).fetchone()
            for entry in level_fixes
        )]
        if not_fixed:
            raise SystemExit(f"這些法術用到「等級修正」寫法，但等級修正表沒有它們：{not_fixed}")

        school_overrides = {}
        for row in rows(book["學派疑義"]):
            keys = []
            for name in split_names(row["建議學派"]):
                if name == "未分類":
                    continue
                if name not in school_names:
                    raise SystemExit(f"學派疑義「{row['法術']}」的學派無效：{name}")
                keys.append(school_names[name])
            for (entry_id,) in db.execute(
                "SELECT e.id FROM spells s JOIN spell_entries e ON e.spell_id=s.id WHERE s.name_zh=? AND s.name_en=? AND e.school=?",
                (row["法術"], row["英文名稱"], row["目前學派文字"] or ""),
            ):
                school_overrides[entry_id] = {"text": row["目前學派文字"] or "", "schools": keys}

    payload = {
        "schools": [{"key": key, "name": name, "pattern": pattern} for key, name, pattern in SCHOOLS],
        "classes": classes,
        "class_aliases": dict(sorted(aliases.items())),
        "dropped_class_aliases": sorted(dropped),
        "level_fixes": dict(sorted(level_fixes.items())),
        "school_overrides": dict(sorted(school_overrides.items())),
    }
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    print(
        f"{OUTPUT.relative_to(ROOT)}: {len(classes)} classes, {len(aliases)} aliases, "
        f"{len(level_fixes)} level fixes, {len(school_overrides)} school overrides"
    )


if __name__ == "__main__":
    main()
